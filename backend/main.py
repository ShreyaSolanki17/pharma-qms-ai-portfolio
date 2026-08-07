import io
import math
import os
import re
from collections import Counter
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional, Union

import fitz  # PyMuPDF
from pypdf import PdfReader
import docx

sys.path.append(str(Path(__file__).resolve().parent.parent))

from fastapi import Depends, FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from backend.agent_pipeline import run_complaint_pipeline, judge_duplicates, run_update_pipeline
from backend import models
from backend.database import SessionLocal, engine, get_db, sync_missing_columns

models.Base.metadata.create_all(bind=engine)
sync_missing_columns()

app = FastAPI(title="Pharma QMS Complaint Management API", version="1.0.0")

# browsers reject allow_credentials=True combined with origins="*", and we don't need cookies anyway
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ComplaintAnalyzeRequest(BaseModel):
    raw_text: str


class ComplaintUpdateRequest(BaseModel):
    correction_text: str
    existing_data: Dict[str, Any]


class ComplaintSaveRequest(BaseModel):
    complaint_ref_no: str
    complaint_source: str
    customer_name: str
    product_name: str
    product_strength: str
    batch_number: str
    manufacturing_date: str
    expiry_date: str
    affected_quantity: Union[str, int]
    originating_block: str
    impacted_npm: str
    complaint_type: str
    complaint_date: str
    complaint_description: str
    initial_severity: str
    priority: str
    risk_assessment: str
    complaint_summary: Optional[str] = ""
    root_cause_recommendation: Optional[str] = ""
    capa_recommendation: Optional[str] = ""
    next_action: Optional[str] = ""


class DuplicateCheckRequest(BaseModel):
    complaint_description: Optional[str] = ""
    product_name: Optional[str] = ""
    batch_number: Optional[str] = ""
    complaint_type: Optional[str] = ""
    top_n: int = 5
    threshold: float = 0.15  # lower = more sensitive

# Duplicate Detection: TF-IDF + Cosine Similarity

_STOP_WORDS = {
    "a", "an", "the", "and", "or", "but", "in", "on", "at", "to", "for",
    "of", "with", "is", "was", "are", "were", "be", "been", "being",
    "have", "has", "had", "do", "does", "did", "will", "would", "shall",
    "should", "may", "might", "can", "could", "this", "that", "these",
    "those", "it", "its", "from", "by", "as", "not", "no", "customer",
    "complaint", "product", "batch", "regarding", "report"
}


def _tokenize(text: str) -> List[str]:
    if not text:
        return []
    tokens = re.findall(r"[a-zA-Z0-9]+", str(text).lower())
    return [t for t in tokens if t not in _STOP_WORDS and len(t) > 1]


def _build_document(complaint) -> str:
    parts = [
        str(complaint.complaint_description or ""),
        str(complaint.product_name or ""),
        str(complaint.batch_number or ""),
        str(complaint.complaint_type or ""),
    ]
    return " ".join([p for p in parts if p])


def _build_query_document(req: "DuplicateCheckRequest") -> str:
    parts = [
        str(req.complaint_description or ""),
        str(req.product_name or ""),
        str(req.batch_number or ""),
        str(req.complaint_type or ""),
    ]
    return " ".join([p for p in parts if p])


def _tf(tokens: List[str]) -> dict:
    count = Counter(tokens)
    total = max(len(tokens), 1)
    return {t: c / total for t, c in count.items()}


def _compute_tfidf_cosine(
    query_doc: str,
    corpus_docs: List[str],
) -> List[float]:
    all_docs = [query_doc] + corpus_docs
    tokenized = [_tokenize(d) for d in all_docs]

    N = len(all_docs)  # includes the query doc itself, not just the corpus
    df: Counter = Counter()
    for tokens in tokenized:
        df.update(set(tokens))

    idf = {term: math.log((N + 1) / (freq + 1)) + 1 for term, freq in df.items()}

    def tfidf_vector(tokens):
        tf = _tf(tokens)
        return {t: tf[t] * idf.get(t, 0) for t in tf}

    vectors = [tfidf_vector(t) for t in tokenized]
    query_vec = vectors[0]

    def cosine(a: dict, b: dict) -> float:
        common = set(a) & set(b)
        dot = sum(a[t] * b[t] for t in common)
        mag_a = math.sqrt(sum(v ** 2 for v in a.values()))
        mag_b = math.sqrt(sum(v ** 2 for v in b.values()))
        if mag_a == 0 or mag_b == 0:
            return 0.0
        return dot / (mag_a * mag_b)

    return [cosine(query_vec, vectors[i + 1]) for i in range(len(corpus_docs))]

def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    extracted_text = ""
    try:
        if ext == ".pdf":
            try:
                pdf_doc = fitz.open(stream=file_bytes, filetype="pdf")
                for page_num, page in enumerate(pdf_doc):
                    if page_num >= 10:
                        break
                    t = page.get_text("text")
                    if t and t.strip():
                        extracted_text += t + "\n"
            except Exception:
                extracted_text = ""

            if not extracted_text.strip():
                pdf_file = io.BytesIO(file_bytes)
                reader = PdfReader(pdf_file)
                for page_num, page in enumerate(reader.pages):
                    if page_num >= 10:
                        break
                    t = page.extract_text()
                    if t:
                        extracted_text += t + "\n"

        elif ext in [".docx", ".doc"]:
            docx_file = io.BytesIO(file_bytes)
            doc = docx.Document(docx_file)
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    extracted_text += paragraph.text.strip() + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        extracted_text += " | ".join(row_cells) + "\n"
        else:
            extracted_text = file_bytes.decode("utf-8", errors="ignore")

    except Exception as err:
        raise HTTPException(status_code=400, detail=f"Failed to process file '{filename}': {str(err)}")
        
    return extracted_text.strip()


@app.get("/")
def read_root():
  return {"status": "online", "system": "Pharma QMS Backend Active"}


@app.post("/api/complaints/analyze")
def analyze_complaint(request: ComplaintAnalyzeRequest):
  try:
    extracted = run_complaint_pipeline(request.raw_text)
    return {"status": "success", "extracted_data": extracted}
  except HTTPException:
    raise
  except Exception as e:
    raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/complaints/update")
def update_complaint(request: ComplaintUpdateRequest):
  """Apply a natural-language correction, returning only the fields that changed."""
  try:
    updated_fields = run_update_pipeline(request.existing_data, request.correction_text)
    return {"status": "success", "updated_fields": updated_fields}
  except HTTPException:
    raise
  except Exception as e:
    raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/complaints/upload")
def upload_complaint_file(file: UploadFile = File(...)):
  try:
    file.file.seek(0)
    contents = file.file.read()
    if not contents:
      raise HTTPException(status_code=400, detail=f"The uploaded file '{file.filename}' is empty (0 bytes).")

    extracted_text = extract_text_from_file_bytes(contents, file.filename)
    if not extracted_text:
      raise HTTPException(
          status_code=400,
          detail=f"Could not extract text from '{file.filename}'. The document may be empty or a scanned image."
      )
    
    extracted_data = run_complaint_pipeline(extracted_text)
    return {
        "status": "success",
        "filename": file.filename,
        "raw_text": extracted_text,
        "extracted_data": extracted_data
    }
  except HTTPException:
    raise
  except Exception as e:
    raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/complaints/check-duplicates")
def check_duplicate_complaint(
    request: DuplicateCheckRequest,
    db: Session = Depends(get_db)
):
  """Find similar existing complaints via TF-IDF, then have the LLM confirm the top matches."""
  try:
    existing = db.query(models.ComplaintModel).all()
    if not existing:
      return {"status": "success", "matches": [], "total_checked": 0}

    corpus_docs = [_build_document(c) for c in existing]
    query_doc   = _build_query_document(request)

    if not query_doc.strip():
      raise HTTPException(status_code=400, detail="complaint_description is required for duplicate checking.")

    scores = _compute_tfidf_cosine(query_doc, corpus_docs)

    # an exact batch match is a near-certain duplicate regardless of text similarity
    query_batch = (request.batch_number or "").strip().lower()

    results = []
    for score, complaint in zip(scores, existing):
      db_batch = (complaint.batch_number or "").strip().lower()
      boosted_score = score
      if query_batch and db_batch and query_batch == db_batch:
        boosted_score = max(score, 0.85)  # exact batch match = almost certainly a duplicate

      if boosted_score >= request.threshold:
        results.append({
            "id": complaint.id,
            "complaint_ref_no": complaint.complaint_ref_no,
            "customer_name": complaint.customer_name,
            "product_name": complaint.product_name,
            "batch_number": complaint.batch_number,
            "complaint_type": complaint.complaint_type,
            "complaint_date": complaint.complaint_date,
            "initial_severity": complaint.initial_severity,
            "complaint_description_snippet": (complaint.complaint_description or "")[:200],
            "similarity_score": round(boosted_score, 4),
            "similarity_pct": round(boosted_score * 100, 1),
        })

    results.sort(key=lambda x: x["similarity_score"], reverse=True)
    top_results = results[: request.top_n]

    # judge on a longer excerpt than the 200-char display snippet - complaint forms often
    # lead with a boilerplate header, and 200 chars can cut off before the actual issue text
    descriptions_by_id = {c.id: (c.complaint_description or "")[:800] for c in existing}
    # let the LLM reason over just the shortlisted matches - catches same-issue-worded-differently
    # cases and rules out coincidental keyword overlap, which cosine similarity alone can't
    verdicts = judge_duplicates(query_doc, [
        {"id": r["id"], "description": descriptions_by_id.get(r["id"], "")} for r in top_results
    ])
    for r in top_results:
      verdict = verdicts.get(r["id"])
      r["ai_verdict"] = verdict.get("is_duplicate") if verdict else None
      r["ai_reasoning"] = verdict.get("reasoning") if verdict else None

    return {
      "status": "success",
      "matches": top_results,
      "total_checked": len(existing),
      "threshold_used": request.threshold,
    }
  except HTTPException:
    raise
  except Exception as e:
    raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/complaints")
async def save_complaint(
    request: ComplaintSaveRequest, db: Session = Depends(get_db)
):
  try:
    db_complaint = models.ComplaintModel(**request.model_dump())
    db.add(db_complaint)
    db.commit()
    db.refresh(db_complaint)
    return {"status": "success", "id": db_complaint.id}
  except IntegrityError:
    db.rollback()
    raise HTTPException(
        status_code=400,
        detail=f"A complaint with Reference Number '{request.complaint_ref_no}' has already been saved."
    )
  except Exception as e:
    db.rollback()
    raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/complaints")
async def get_all_complaints(db: Session = Depends(get_db)):
  try:
    complaints = db.query(models.ComplaintModel).all()
    return complaints
  except Exception as e:
    raise HTTPException(status_code=500, detail=str(e))


@app.delete("/api/complaints/{complaint_id}")
async def delete_complaint(complaint_id: int, db: Session = Depends(get_db)):
  try:
    record = db.query(models.ComplaintModel).filter(models.ComplaintModel.id == complaint_id).first()
    if not record:
      raise HTTPException(status_code=404, detail="Complaint record not found.")
    
    db.delete(record)
    db.commit()
    return {"status": "success", "message": f"Complaint ID {complaint_id} deleted successfully."}
  except Exception as e:
    db.rollback()
    raise HTTPException(status_code=500, detail=str(e))